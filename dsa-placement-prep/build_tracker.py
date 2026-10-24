"""Builds DSA_Progress_Tracker.xlsx from problems.py. Run: python3 build_tracker.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

from problems import flat

rows = flat()
N_EXTRA = 120  # blank rows for problems you add yourself
LAST = 1 + len(rows) + N_EXTRA

HEAD = PatternFill("solid", fgColor="1F3864")
HF = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF9E5")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

STATUSES = ["Not started", "Attempted", "Solved independently",
            "Solved with hints", "Solved after viewing solution"]

wb = Workbook()

# ---------------------------------------------------------------- Tracker
ws = wb.active
ws.title = "Tracker"
cols = [
    ("#", 5), ("Problem title", 44), ("Platform", 12), ("Problem URL", 14),
    ("Topic", 26), ("Difficulty", 10), ("Priority", 12), ("Week", 6),
    ("Date first attempted", 14), ("Status", 28), ("Solved independently?", 12),
    ("Attempts", 9), ("Time taken (min)", 10), ("Time complexity", 14),
    ("Space complexity", 14), ("Key lesson learned", 44),
    ("Accepted on platform? (Y/N)", 13), ("Date solved", 12),
    ("Next revision due", 13), ("Revisions done", 10),
    ("Problem key (edit to match same problem across platforms)", 34),
    ("Counts as unique solve (auto)", 12), ("Duplicate warning (auto)", 14),
]
for i, (name, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=i, value=name)
    c.fill, c.font, c.border = HEAD, HF, BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 48
ws.freeze_panes = "C2"

for r in range(2, LAST + 1):
    idx = r - 2
    p = rows[idx] if idx < len(rows) else None
    if p:
        ws.cell(r, 1, p["n"])
        ws.cell(r, 2, p["title"])
        ws.cell(r, 3, p["platform"])
        c = ws.cell(r, 4, "Open")
        c.hyperlink = p["url"]
        c.font = Font(color="0563C1", underline="single")
        ws.cell(r, 5, p["topic"])
        ws.cell(r, 6, p["diff"])
        ws.cell(r, 7, p["prio"])
        ws.cell(r, 8, p["week"])
        ws.cell(r, 10, "Not started")
    # auto formulas (also on blank rows so added problems work)
    ws.cell(r, 11, f'=IF(J{r}="Solved independently","Yes",IF(OR(J{r}="",J{r}="Not started"),"","No"))')
    ws.cell(r, 19, f'=IF(R{r}="","",R{r}+IF(T{r}=0,3,IF(T{r}=1,7,IF(T{r}=2,21,45))))')
    ws.cell(r, 21, f'=IF(B{r}="","",LOWER(TRIM(B{r})))')
    solved = (f'OR(J{r}="Solved independently",J{r}="Solved with hints",'
              f'J{r}="Solved after viewing solution")')
    ws.cell(r, 22, f'=IF(AND(U{r}<>"",{solved},COUNTIFS($U$2:U{r-1},U{r},$V$2:V{r-1},1)=0),1,0)'
            if r > 2 else f'=IF(AND(U{r}<>"",{solved}),1,0)')
    ws.cell(r, 23, f'=IF(AND(U{r}<>"",COUNTIF($U$2:$U${LAST},U{r})>1),"DUPLICATE KEY","")')
    for col in range(1, len(cols) + 1):
        cell = ws.cell(r, col)
        cell.border = BORDER
        if col in (9, 10, 12, 13, 14, 15, 16, 17, 18, 20):
            cell.fill = INPUT
    for col in (9, 18):
        ws.cell(r, col).number_format = "yyyy-mm-dd"
    ws.cell(r, 19).number_format = "yyyy-mm-dd"
    ws.cell(r, 16).alignment = Alignment(wrap_text=True, vertical="top")
    if not p:
        for col in (2, 3, 5, 6, 7, 10):
            ws.cell(r, col).fill = INPUT
        ws.cell(r, 10, "Not started")

ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{LAST}"

def dv(formula, col):
    d = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(d)
    d.add(f"{col}2:{col}{LAST}")

dv('"' + ",".join(STATUSES) + '"', "J")
dv('"LeetCode,HackerRank,GeeksforGeeks,CodeChef,Other"', "C")
dv('"Easy,Medium,Hard"', "F")
dv('"Essential,Recommended,Optional"', "G")
dv('"Y,N"', "Q")

rng = f"J2:J{LAST}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Solved independently"'], fill=PatternFill("solid", bgColor="C6EFCE")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Solved with hints"'], fill=PatternFill("solid", bgColor="FFEB9C")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Solved after viewing solution"'], fill=PatternFill("solid", bgColor="F8CBAD")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Attempted"'], fill=PatternFill("solid", bgColor="DDEBF7")))
ws.conditional_formatting.add(f"S2:S{LAST}", FormulaRule(formula=['AND(S2<>"",S2<=TODAY())'], fill=PatternFill("solid", bgColor="FFC7CE")))
ws.conditional_formatting.add(f"W2:W{LAST}", CellIsRule(operator="equal", formula=['"DUPLICATE KEY"'], font=Font(color="9C0006", bold=True)))

# -------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
db.column_dimensions["A"].width = 38
for c in "BCDEF":
    db.column_dimensions[c].width = 16
db["A1"] = "DSA progress dashboard (unique problems only)"
db["A1"].font = Font(bold=True, size=14)
db["A2"] = ("Counts come only from the Tracker sheet. A problem counts once, "
            "when its status is solved (any of the 3 solved kinds) and its key has not already been counted.")
db["A2"].alignment = Alignment(wrap_text=True)
db.merge_cells("A2:F2")
db.row_dimensions[2].height = 32

T = lambda col: f"Tracker!${col}$2:${col}${LAST}"
r = 4
def head(txt, cols_):
    global r
    for i, t in enumerate([txt] + cols_):
        c = db.cell(r, 1 + i, t)
        c.fill, c.font = HEAD, HF
    r += 1

head("Overall", ["Count"])
db.cell(r, 1, "Unique problems solved (any way)"); db.cell(r, 2, f"=SUM({T('V')})"); total_row = r; r += 1
db.cell(r, 1, "  Solved independently"); db.cell(r, 2, f'=COUNTIFS({T("V")},1,{T("J")},"Solved independently")'); indep_row = r; r += 1
db.cell(r, 1, "  Solved with hints"); db.cell(r, 2, f'=COUNTIFS({T("V")},1,{T("J")},"Solved with hints")'); r += 1
db.cell(r, 1, "  Solved after viewing solution"); db.cell(r, 2, f'=COUNTIFS({T("V")},1,{T("J")},"Solved after viewing solution")'); r += 1
db.cell(r, 1, "Attempted but not solved"); db.cell(r, 2, f'=COUNTIF({T("J")},"Attempted")'); r += 1
db.cell(r, 1, "Independence rate"); db.cell(r, 2, f"=IF(B{total_row}=0,0,B{indep_row}/B{total_row})"); db.cell(r, 2).number_format = "0%"; r += 1
db.cell(r, 1, "Revisions overdue"); db.cell(r, 2, f'=COUNTIFS({T("S")},"<="&TODAY(),{T("S")},"<>")'); r += 1
db.cell(r, 1, "Rows with duplicate keys (fix these)"); db.cell(r, 2, f'=COUNTIF({T("W")},"DUPLICATE KEY")'); r += 2

head("By platform (unique solves)", ["Solved", "Independent"])
for plat in ["LeetCode", "HackerRank", "GeeksforGeeks", "CodeChef", "Other"]:
    db.cell(r, 1, plat)
    db.cell(r, 2, f'=SUMIFS({T("V")},{T("C")},"{plat}")')
    db.cell(r, 3, f'=COUNTIFS({T("V")},1,{T("C")},"{plat}",{T("J")},"Solved independently")')
    r += 1
r += 1
head("By difficulty (unique solves)", ["Solved", "Independent"])
for d in ["Easy", "Medium", "Hard"]:
    db.cell(r, 1, d)
    db.cell(r, 2, f'=SUMIFS({T("V")},{T("F")},"{d}")')
    db.cell(r, 3, f'=COUNTIFS({T("V")},1,{T("F")},"{d}",{T("J")},"Solved independently")')
    r += 1
r += 1
head("By topic (unique solves)", ["Solved", "Listed", "Independent"])
topics = []
for p in rows:
    if p["topic"] not in topics:
        topics.append(p["topic"])
for t in topics:
    db.cell(r, 1, t)
    db.cell(r, 2, f'=SUMIFS({T("V")},{T("E")},A{r})')
    db.cell(r, 3, f'=COUNTIF({T("E")},A{r})')
    db.cell(r, 4, f'=COUNTIFS({T("V")},1,{T("E")},A{r},{T("J")},"Solved independently")')
    r += 1
r += 1
head("Milestones", ["Target", "Reached?"])
for m in [25, 50, 100, 150, 200]:
    db.cell(r, 1, f"Milestone: {m} unique solved")
    db.cell(r, 2, m)
    db.cell(r, 3, f'=IF(B{total_row}>={m},"Yes","Not yet")')
    r += 1

# -------------------------------------------------------- Evidence / mocks
ev = wb.create_sheet("Verification log")
for i, (n, w) in enumerate([("Platform", 16), ("Profile URL", 40), ("Date checked", 14),
                            ("Platform-reported solved count", 16), ("My tracker's unique count", 16),
                            ("Screenshot/PDF saved? (Y/N)", 14), ("Notes (e.g. counts include premium/duplicates?)", 50)], 1):
    c = ev.cell(1, i, n); c.fill, c.font = HEAD, HF
    c.alignment = Alignment(wrap_text=True)
    ev.column_dimensions[get_column_letter(i)].width = w
for i, plat in enumerate(["LeetCode", "HackerRank", "GeeksforGeeks", "CodeChef"], 2):
    ev.cell(i, 1, plat)
ev["A8"] = ("Update monthly. Quote platform counts as 'platform-reported'. Quote your tracker count "
            "as your own record. Only put a number on a resume that you have just confirmed on your live profile.")
ev["A8"].alignment = Alignment(wrap_text=True); ev.merge_cells("A8:G8"); ev.row_dimensions[8].height = 32

mk = wb.create_sheet("Mock log")
for i, (n, w) in enumerate([("Date", 12), ("Type (OA / timed set / interview)", 22), ("Source", 24),
                            ("Problems given", 10), ("Fully solved (all tests)", 12), ("Partially solved", 10),
                            ("Time used / limit (min)", 14), ("Where I lost time", 40), ("Weak topics exposed", 36),
                            ("Action for next week", 40)], 1):
    c = mk.cell(1, i, n); c.fill, c.font = HEAD, HF
    c.alignment = Alignment(wrap_text=True)
    mk.column_dimensions[get_column_letter(i)].width = w

wk = wb.create_sheet("Weekly review")
for i, (n, w) in enumerate([("Week", 7), ("Week start", 12), ("New problems solved", 12), ("Revisions done", 10),
                            ("Solved independently (of new)", 14), ("Avg minutes per Easy", 12),
                            ("Avg minutes per Medium", 12), ("Strongest topic", 22), ("Weakest topic", 22),
                            ("Decision: advance / repeat / add problems", 28), ("Change for next week", 44)], 1):
    c = wk.cell(1, i, n); c.fill, c.font = HEAD, HF
    c.alignment = Alignment(wrap_text=True)
    wk.column_dimensions[get_column_letter(i)].width = w
for w_ in range(1, 15):
    wk.cell(1 + w_, 1, w_)

# ------------------------------------------------------------ Problem list
pl = wb.create_sheet("Problem list")
hdr = ["Order", "Week", "Topic", "Problem", "Platform", "Link", "Difficulty",
       "Concept tested", "Prerequisites", "Priority"]
for i, h in enumerate(hdr, 1):
    c = pl.cell(1, i, h); c.fill, c.font = HEAD, HF
for rr, p in enumerate(rows, 2):
    vals = [p["n"], p["week"], p["topic"], p["title"], p["platform"], p["url"], p["diff"],
            p["concept"], p["prereq"], p["prio"]]
    for i, v in enumerate(vals, 1):
        pl.cell(rr, i, v)
    pl.cell(rr, 6).hyperlink = p["url"]
    pl.cell(rr, 6).font = Font(color="0563C1", underline="single")
for i, w in enumerate([7, 6, 32, 46, 12, 70, 10, 38, 28, 13], 1):
    pl.column_dimensions[get_column_letter(i)].width = w
pl.freeze_panes = "A2"
pl.auto_filter.ref = f"A1:J{len(rows)+1}"

wb.save("DSA_Progress_Tracker.xlsx")
print("saved", len(rows), "problems")
