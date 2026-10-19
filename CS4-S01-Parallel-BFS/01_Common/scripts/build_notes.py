"""Content of the two individual technical notes (extended versions, about 5-6 pages each).

Imported by build_docs.py. Every number comes from the CSV files (raw or derived) through the helpers of build_docs.py.
"""
import math
import statistics as st

from docx.shared import Inches, Pt

from build_docs import (D, ENV, FAM, NDENS, NMAX, NMIN, OUT_N1, OUT_N2, RAW, ROOT, SYNC4, TMAX, cv_hi, EQ_COMPLEX, EQ_EFF, EQ_SPEEDUP, EQ_TEPS, EQ_TEPS_SEQ,
                        bullets, cv_max, dmt, ef, equation, f1, f2, figure_row, grid_par_us, grid_seq_us, mt, ms, n_cfg, n_noisy, new_doc, p_lev_us,
                        para, r, seq_level_ns, sig, sp, table, work_per_level, worst_ratio, _frac, _r, _sub)

SIZES = sorted(D.Problem_Size.unique())
N_SEQ_RUNS = int((RAW.Algorithm == "SEQ_BFS").sum())
N_PAR_RUNS = int((RAW.Algorithm != "SEQ_BFS").sum())
N_SEQ_CFG = int((D.Algorithm == "SEQ_BFS").sum())
N_PAR_CFG = int((D.Algorithm != "SEQ_BFS").sum())
N_STRUCT_RUNS = int(RAW[(RAW.Algorithm == "SEQ_BFS") & RAW.Dataset.isin([d for d, _ in FAM])].shape[0])
LOG2 = lambda n: int(round(math.log2(n)))
SZ = lambda n: f"2^{LOG2(n)}"
MINUS = "−"


def code_lines(doc, lines, size=8):
    """monospace block (pseudocode)"""
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(ln)
        run.font.name = "Consolas"
        run.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def seq_runs(ds, n):
    return [float(x) for x in RAW[(RAW.Dataset == ds) & (RAW.Problem_Size == n) & (RAW.Algorithm == "SEQ_BFS")].Execution_Time]


# ============================================================================== NOTE 1
def trace_two_routes():
    """BFS trace on tests/cases/two_routes.graph, simulated here and compared with the hand-written expected file"""
    lines = [l for l in (ROOT / "tests/cases/two_routes.graph").read_text().splitlines() if l and not l.startswith("#")]
    n, m = map(int, lines[0].split())
    adj = {i: [] for i in range(n)}
    for l in lines[1:]:
        u, v = map(int, l.split())
        adj[u].append(v)
        adj[v].append(u)
    for k in adj:
        adj[k].sort()
    dist = [-1] * n
    dist[0] = 0
    q = [0]
    head = 0
    rows = [["0", "-", "-", "[0]", " ".join(str(x) for x in dist)]]
    step = 0
    while head < len(q):
        u = q[head]
        head += 1
        step += 1
        new = []
        for v in adj[u]:
            if dist[v] == -1:
                dist[v] = dist[u] + 1
                q.append(v)
                new.append(v)
        rows.append([str(step), str(u), ", ".join(map(str, new)) if new else "none", "[" + ", ".join(map(str, q[head:])) + "]", " ".join(str(x) for x in dist)])
    exp = [int(x) for x in (ROOT / "tests/cases/two_routes.expected").read_text().split()[2:]]
    assert dist == exp, "trace disagrees with the hand-written expected file"
    return rows, n, m


CASE_NOTES = {
    "path5": "chain; distances grow by one", "path5_mid": "source in the middle (two directions)", "star6": "hub with many neighbours, leaf source",
    "cycle6": "two routes around a cycle; shorter wins", "grid3x3": "2-D structure; distance = row + column", "disconnected": "unreachable vertices must stay -1",
    "isolated_source": "source without any neighbour", "complete5": "dense graph; everything at distance 1", "single": "smallest possible graph (n = 1)",
    "binary_tree7": "branching; levels 0, 1, 2", "dup_selfloop": "input cleaning: duplicate edge and self-loop", "two_routes": "long and short route; must choose the short one"}


def case_rows():
    rows = []
    for g in sorted((ROOT / "tests/cases").glob("*.graph")):
        name = g.stem
        h = [l for l in g.read_text().splitlines() if l and not l.startswith("#")][0].split()
        e = (ROOT / "tests/cases" / f"{name}.expected").read_text().split()
        src, dist = e[1], " ".join(e[2:])
        rows.append([name, h[0], h[1], src, dist, CASE_NOTES[name]])
    return rows


def build_note1():
    doc = new_doc(size=10.5)
    eqc = [0]
    para(doc, "Individual Technical Note — Student 1", bold=True, align="center", size=14, after=2)
    para(doc, "Sequential BFS, graph preparation, correctness verification and baseline measurements", align="center", size=11, after=2)
    para(doc, "CS4-S01 Parallel Breadth-First Search — BCSE412L Parallel Computing | [Student 1 name, Reg. No.] | Batch [ID]", align="center", size=9.5, after=6)

    doc.add_heading("1  My contribution", level=1)
    para(doc, "I was responsible for the baseline of the case study: explaining the BFS algorithm, implementing the sequential BFS in C++, preparing reproducible graph inputs, verifying correctness, "
              "measuring the sequential baseline and documenting the baseline methodology. My partner built the OpenMP version on top of the same graph code and conventions, so that every speedup is measured against my baseline. "
              "Both of us wrote the literature survey and the integrated report. This note describes my part in more detail than the integrated report can.", align="justify")
    table(doc, ["Item", "Where"], [
        ["Sequential BFS and independent validator", "02_Lucky_Sequential_BFS/bfs_seq.hpp"],
        ["Baseline benchmark program", "02_Lucky_Sequential_BFS/bfs_seq_main.cpp"],
        ["CSR graph, four generators, graph-file loader, start-vertex rule, graph checksum", "01_Common/graph.hpp"],
        ["12 hand-computed test graphs with expected distances", "01_Common/tests/cases/*.graph, *.expected"],
        ["Baseline measurements", "01_Common/data/raw/bfs_raw.csv (rows with Algorithm = SEQ_BFS)"],
    ], widths=[3.6, 2.9], size=8.5, caption="Table 1. My files and results.")

    doc.add_heading("2  The BFS algorithm and its cost", level=1)
    para(doc, "BFS starts at a source vertex s with distance 0. A queue holds the vertices that have been discovered but not yet expanded. The algorithm removes the vertex u at the head of the queue and, "
              "for every neighbour v of u that has not been discovered (dist[v] = −1), sets dist[v] = dist[u] + 1 and appends v. Because the queue is first-in-first-out, vertices are expanded in order of distance, "
              "so the first time a vertex is reached is along a shortest path (in number of edges). Vertices that are never reached keep the value −1.", align="justify")
    para(doc, "Pseudocode:  dist[*] = −1; dist[s] = 0; push(s);  while queue not empty: u = pop(); for each neighbour v of u: if dist[v] == −1: dist[v] = dist[u] + 1; push(v).", italic=True, size=9.5)
    rows, n6, m6 = trace_two_routes()
    para(doc, f"**A worked example.** Table 2 traces the algorithm on the test graph two_routes ({n6} vertices, {m6} edges: a long route 0-1-2-3-4 and a short route 0-5-4), starting at vertex 0. "
              "Vertex 4 is first discovered from vertex 5 at distance 2, not from vertex 3 at distance 4, which shows why the queue order gives shortest distances. The final distances agree with the hand-written expected file of the test suite.", align="justify")
    table(doc, ["Step", "Removed", "Newly discovered", "Queue afterwards", "dist[0..5]"], rows, widths=[0.5, 0.7, 1.4, 1.7, 1.9], size=8.5,
          caption="Table 2. Trace of sequential BFS on the graph two_routes (generated by simulating the algorithm; compared with tests/cases/two_routes.expected).")
    para(doc, "**Complexity.** Each vertex is pushed and popped once and each adjacency entry is read once, so the running time and memory are", after=1, align="justify")
    equation(doc, EQ_COMPLEX, eqc)
    para(doc, "for n vertices and m edges, with O(n) extra memory for the distance array and the queue. In the code (02_Lucky_Sequential_BFS/bfs_seq.hpp) the queue is a plain integer array of length n, which is enough because no vertex is pushed twice. "
              "In our setting each undirected edge appears twice in the adjacency lists, so the loop reads 2m entries. We define the throughput of a run in traversed edges per second (TEPS) with the Graph500 convention", align="justify", after=1)
    equation(doc, EQ_TEPS_SEQ, eqc)
    para(doc, "where m_c is the number of undirected edges in the component reached from the start vertex.", align="justify")

    doc.add_heading("3  Graph representation and preparation", level=1)
    para(doc, "All graphs are undirected, simple and stored in compressed sparse row (CSR) form: an offset array off[0..n] and a neighbour array adj, so that the neighbours of u are adj[off[u]..off[u+1]−1]. "
              "Compared with a vector of vectors, CSR needs only two allocations, stores the neighbours of a vertex next to each other in memory, and is easy to share between threads read-only. "
              "Building the CSR removes self-loops and duplicate edges and sorts each neighbour list. I wrote four generators in 01_Common/graph.hpp:", align="justify")
    bullets(doc, [
        "**Path** (n vertices in a line): n levels, one vertex per level; the worst case for parallelism.",
        "**Grid** (√n × √n, four neighbours): about 2√n levels; frontier width up to √n.",
        "**Random** (uniform random pairs, average degree 16 or other values): few levels, very wide frontiers.",
        "**R-MAT** (Graph500-like parameters 0.57 / 0.19 / 0.19 / 0.05, vertex labels randomly permuted): skewed degrees with hubs.",
    ])
    para(doc, "The generators use our own SplitMix64 random generator with a fixed seed (12345), so the same graph is produced on every compiler and operating system. "
              "Each CSV row stores a checksum of the graph; the analysis script stops if the sequential and parallel rows of one graph have different checksums. "
              "The start vertex is vertex 0 for the path and grid, and the lowest-numbered vertex of highest degree for random and R-MAT graphs, so that these traversals start inside the large connected part. "
              "Graph construction is never part of the timed region. Table 3 shows what the generated graphs look like.", align="justify")
    rows = []
    for ds, lab in FAM:
        for n in SIZES:
            s = r(ds, n, 0)
            rows.append([lab, SZ(n), f"{int(s.Edges_Undirected):,}", f"{s.Avg_Degree:.2f}", f"{int(s.Max_Degree):,}", int(s.Levels), f"{100 * s.Reached_Vertices / s.Vertices:.1f}"])
    table(doc, ["Graph", "n", "Edges", "Avg. degree", "Max degree", "Levels", "Reached (% of n)"], rows, widths=[1.3, 0.55, 1.0, 0.85, 0.85, 0.7, 1.0], size=8,
          caption="Table 3. Structure of the generated graphs as seen from the start vertex (from data/derived/bfs_derived.csv).")
    rs = r("rmat_ef8", NMAX, 0)
    para(doc, f"Three observations follow from Table 3. The path and the grid have regular degrees (at most 2 and 4), while the R-MAT graph has hubs (maximum degree {int(rs.Max_Degree):,} at the largest size). "
              f"The number of levels differs enormously for the same number of vertices: {int(r('path', NMAX, 0).Levels):,} for the path, {int(r('grid', NMAX, 0).Levels):,} for the grid and only {int(r('random_d16', NMAX, 0).Levels)}–{int(rs.Levels)} for the random and R-MAT graphs. "
              f"Finally, the R-MAT graph is not fully connected from the start vertex: only {100 * rs.Reached_Vertices / rs.Vertices:.0f}% of its vertices are reached, because the skewed generator leaves many isolated vertices. "
              "This is why throughput is always computed from the edges of the reached component and not from all edges.", align="justify")

    doc.add_heading("4  Correctness verification", level=1)
    para(doc, "I verified the sequential BFS in three ways (log: 01_Common/logs/correctness_test_output.txt). First, twelve small graphs have expected distances that I worked out by hand (Table 4). "
              "The program compares the result of the sequential BFS and of the parallel BFS (five thread counts and three loop schedules) with these expected distances.", align="justify")
    table(doc, ["Test graph", "n", "Edges", "Source", "Expected distances (−1 = unreachable)", "What it tests"], case_rows(), widths=[1.05, 0.3, 0.45, 0.5, 1.9, 2.3], size=7.5,
          caption="Table 4. The twelve hand-computed test cases (tests/cases).")
    bullets(doc, [
        "**Independent validator.** validate_bfs() checks that dist[s] = 0, that two neighbouring vertices differ by at most one level, that a reachable and an unreachable vertex are never adjacent, and that every reached vertex has a neighbour one level closer. "
        "It does not reuse the BFS code, so a bug in the BFS cannot hide itself. Together these rules characterise exactly the valid BFS distance arrays.",
        "**Randomized cross-checks.** Sequential and parallel results are compared on generated path, grid, random, R-MAT and disconnected graphs, with 1, 2, 4 and 7 threads and static, dynamic and guided schedules, repeated many times to expose rare races. Error handling (bad source, bad edge) is tested as well.",
        "**Counting.** The test program runs 1,402 checks in total and all pass. By its structure, part 1 (the twelve cases above) contributes 205 checks (12 × (1 + 1 + 15) + 1) and part 2 the remaining 1,197.",
        "**During the experiments.** Every one of the 588 timed runs was also compared with the sequential distances and stored with Verified = 1; the analysis script refuses to continue if a single run failed.",
    ])

    doc.add_heading("5  Baseline measurement methodology", level=1)
    para(doc, "The timed region is a single call of bfs_seq(), which includes allocating and initialising its dist and queue arrays but excludes graph generation, validation and statistics. "
              "The clock is std::chrono::steady_clock, in seconds. For every graph one warm-up run is discarded, then seven timed runs are recorded as separate rows of the raw CSV with the columns required by the manual. "
              "I report the median because a few slow runs (caused by other processes on a shared virtual machine) would distort a mean; the mean, standard deviation, coefficient of variation (CV), minimum and maximum are kept in the derived CSV. "
              f"The machine is a shared cloud virtual machine ({ENV['cpu_model']}, {ENV['logical_cpus']} logical CPUs, {ENV['ram_gb']} GB RAM, {ENV['compiler']}, flags -O3 -std=c++17 -fopenmp). "
              "The sequential time of each graph is the baseline T_seq in the speedup", align="justify", after=1)
    equation(doc, EQ_SPEEDUP, eqc)
    para(doc, "of the parallel program, so the quality of this baseline directly decides how good the parallel results look. That is why I measured it carefully and why Section 8 tests how much the choice of the summary statistic matters.", align="justify")

    doc.add_heading("6  Baseline results", level=1)
    rows = []
    for ds, lab in FAM:
        for n in SIZES:
            s = r(ds, n, 0)
            rows.append([lab, SZ(n), f"{int(s.Edges_Undirected):,}", int(s.Levels), f"{s.Median * 1e3:.3f}", f"{s.CV_percent:.0f}", f"{s.MTEPS:.0f}", f"{s.Median / (s.Vertices + s.Edges_Undirected) * 1e9:.1f}"])
    table(doc, ["Graph", "n", "Edges", "Levels", "Median time (ms)", "CV (%)", "MTEPS", "ns per (n+m)"], rows, widths=[1.3, 0.5, 0.95, 0.65, 1.0, 0.55, 0.6, 0.85], size=8,
          caption="Table 5. Sequential BFS baseline (7 runs per row; 01_Common/data/derived/bfs_derived.csv). The last column is the median time divided by n + m, in nanoseconds.")
    figure_row(doc, [("fig8_baseline_scaling.png", "Figure 1. Sequential BFS: median time (left) and throughput (right) versus number of vertices.")], width=6.3)
    grow = {ds: ms(ds, NMAX, 0) / ms(ds, 2 ** 18, 0) for ds, _ in FAM}
    para(doc, f"If the cost were exactly proportional to n + m, the last column of Table 5 would be constant. For the path it is: {r('path', NMIN, 0).Median / (r('path', NMIN, 0).Vertices + r('path', NMIN, 0).Edges_Undirected) * 1e9:.1f} ns at 2^14 and "
              f"{r('path', NMAX, 0).Median / (r('path', NMAX, 0).Vertices + r('path', NMAX, 0).Edges_Undirected) * 1e9:.1f} ns at 2^20, because a path touches memory in order. "
              f"For the random graph the cost per element rises from {r('random_d16', 2**16, 0).Median / (r('random_d16', 2**16, 0).Vertices + r('random_d16', 2**16, 0).Edges_Undirected) * 1e9:.1f} ns at 2^16 to "
              f"{r('random_d16', NMAX, 0).Median / (r('random_d16', NMAX, 0).Vertices + r('random_d16', NMAX, 0).Edges_Undirected) * 1e9:.1f} ns at 2^20, and its throughput falls from {mt('random_d16', 2**18, 0):.0f} to {mt('random_d16', NMAX, 0):.0f} MTEPS. "
              f"Going from 2^18 to 2^20 vertices multiplies the time by {grow['path']:.1f} (path), {grow['grid']:.1f} (grid), {grow['random_d16']:.1f} (random) and {grow['rmat_ef8']:.1f} (R-MAT), where 4 would mean a perfectly linear algorithm. "
              "The theory says O(n + m), so the extra cost must come from the machine, not from the algorithm: random neighbour lists make the BFS jump to unpredictable positions of the distance array, and once the graph is larger than the fast caches most of these accesses wait for main memory. "
              "I did not measure cache events, so this is a likely explanation, not a proven one.", align="justify")

    doc.add_heading("7  Run-to-run variation", level=1)
    figure_row(doc, [("fig9_baseline_run_variation.png", "Figure 2. The seven sequential runs of every graph, each divided by the median of its configuration (raw data).")], width=6.3)
    seq_cv = D[D.Algorithm == "SEQ_BFS"].CV_percent
    sub_ms = D[(D.Algorithm == "SEQ_BFS") & (D.Median < 1e-3)]
    big = D[(D.Algorithm == "SEQ_BFS") & (D.Median >= 1e-3)]
    _sq = RAW[RAW.Algorithm == "SEQ_BFS"].copy()
    _sq["rel"] = _sq.Execution_Time / _sq.groupby(["Dataset", "Problem_Size"]).Execution_Time.transform("median")
    _w = _sq.loc[_sq.rel.idxmax()]
    mx = float(_w.rel)
    para(doc, f"Figure 2 shows the {N_STRUCT_RUNS} sequential runs of the 16 structure and size configurations without removing any of them (the other {N_SEQ_RUNS - N_STRUCT_RUNS} runs belong to the density sweep and are in the CSV). "
              f"Over all {N_SEQ_CFG} baseline configurations the CV is between {seq_cv.min():.0f}% and {seq_cv.max():.0f}%; "
              f"it averages {sub_ms.CV_percent.mean():.0f}% for the {len(sub_ms)} baselines that take less than one millisecond and {big.CV_percent.mean():.0f}% for the {len(big)} longer ones. "
              f"The largest outlier is a run that took {mx:.2f} times the median ({_w.Dataset}, {int(_w.Problem_Size):,} vertices). Short runs are noisy because a single interruption by the operating system or by another program on the shared machine is large compared with the run itself. "
              "I kept these runs and used the median, which is not affected by one slow run. A consequence for the integrated report is that speedups of the smallest graphs have a wide uncertainty, which the report states explicitly.", align="justify")

    doc.add_heading("8  How sensitive are the speedups to my baseline?", level=1)
    rows = []
    for ds, lab in FAM:
        t = seq_runs(ds, NMAX)
        par4 = r(ds, NMAX, TMAX).Median
        rows.append([lab, f"{st.median(t) * 1e3:.2f}", f"{st.mean(t) * 1e3:.2f}", f"{min(t) * 1e3:.2f}", f"{st.median(t) / par4:.3f}", f"{st.mean(t) / par4:.3f}", f"{min(t) / par4:.3f}"])
    table(doc, ["Graph (n = 2^20)", "Median baseline (ms)", "Mean baseline (ms)", "Fastest baseline (ms)", "Speedup, median baseline", "Speedup, mean baseline", "Speedup, fastest baseline"], rows,
          widths=[1.2, 0.85, 0.85, 0.85, 0.95, 0.95, 0.95], size=8, caption=f"Table 6. Speedup with {TMAX} threads if the baseline is summarised by the median (used in the report), the mean, or the fastest run. The parallel time is always the median.")
    para(doc, f"The choice of summary statistic changes the speedups only slightly at the largest size (for the random graph {st.median(seq_runs('random_d16', NMAX)) / r('random_d16', NMAX, TMAX).Median:.2f} with the median and "
              f"{min(seq_runs('random_d16', NMAX)) / r('random_d16', NMAX, TMAX).Median:.2f} with the fastest run), so the conclusions of the report do not depend on it. The strictest, and for the parallel program least flattering, choice would be the fastest sequential run, "
              "because it is the best the baseline can do; the median is a fairer typical value. A different kind of sensitivity is not tested here: a faster sequential algorithm (for example direction-optimizing BFS) would lower all speedups.", align="justify")

    doc.add_heading("9  Limitations and what I would improve", level=1)
    bullets(doc, [
        "The baseline was measured on one shared cloud machine; neighbouring workloads may disturb timings, and I cannot tell whether the four logical CPUs are physical cores.",
        "All graphs are synthetic and only one start vertex is used per graph; real graphs (SNAP, SuiteSparse) have community structure and locality that would change the cost per edge.",
        "The cache explanation of the rising time per edge is not measured; hardware counters (for example with perf on Linux) would confirm or reject it.",
        "The queue-based algorithm is the standard textbook baseline, not the fastest possible sequential BFS.",
        "Windows compilation of the programs was not tested; the tested route is Linux (WSL on Windows).",
    ])
    para(doc, f"**Summary of my contribution.** I implemented and validated a sequential BFS, built the graph generators and the test suite that both of us rely on, measured {N_SEQ_RUNS} baseline runs on {N_SEQ_CFG} graphs, and documented the methodology and the uncertainty of the baseline. "
              "All numbers in this note come from data/raw/bfs_raw.csv and data/derived/bfs_derived.csv.", align="justify")
    doc.save(OUT_N1 / "Technical_Note_Student1.docx")


# ============================================================================== NOTE 2
def exact_speedup(ds, n, t):
    """speedup from the RAW medians (the derived CSV stores only 6 significant digits)"""
    base = RAW[(RAW.Dataset == ds) & (RAW.Problem_Size == n) & (RAW.Algorithm == "SEQ_BFS")].Execution_Time.median()
    par = RAW[(RAW.Dataset == ds) & (RAW.Problem_Size == n) & (RAW.Algorithm != "SEQ_BFS") & (RAW.Threads == t)].Execution_Time.median()
    return base / par


def karp_flatt(s, p):
    return (1 / s - 1 / p) / (1 - 1 / p)


def build_note2():
    doc = new_doc(size=10.5)
    eqc = [0]
    para(doc, "Individual Technical Note — Student 2", bold=True, align="center", size=14, after=2)
    para(doc, "OpenMP parallel BFS, thread scaling, speedup, efficiency, throughput and synchronization", align="center", size=11, after=2)
    para(doc, "CS4-S01 Parallel Breadth-First Search — BCSE412L Parallel Computing | [Student 2 name, Reg. No.] | Batch [ID]", align="center", size=9.5, after=6)

    doc.add_heading("1  My contribution", level=1)
    para(doc, "I implemented the level-synchronous parallel BFS in C++ with OpenMP (03_Girlfriend_OpenMP_BFS/bfs_par.hpp and bfs_par_main.cpp), ran it with several thread counts, measured execution time, speedup, efficiency and traversal throughput, "
              "and analysed synchronization, parallel overhead and scalability. I used the graph code, start-vertex rule, timing convention and correctness checks prepared by Student 1, so that the two programs are directly comparable and every speedup has an identified baseline.", align="justify")

    doc.add_heading("2  Parallel algorithm and frontier processing", level=1)
    para(doc, "BFS is parallelized level by level. The frontier (the vertices at distance L) is shared among threads; the next frontier (distance L+1) is built while this happens. "
              "The whole traversal is inside one OpenMP parallel region, so threads are created once and not once per level. The steps of one level are:", align="justify", after=2)
    code_lines(doc, [
        "parallel region with T threads:",
        "  parallel for v in 0..n-1:   dist[v] = -1",
        "  one thread:                 dist[s] = 0; frontier = [s]; level = 0",
        "  repeat",
        "     buf[tid] = empty                                   // private next-frontier buffer",
        "     parallel for u in frontier (schedule dynamic,64):  // EXPAND",
        "        for each neighbour v of u:",
        "           if dist[v] == -1 and CAS(dist[v], -1 -> level+1) succeeds:",
        "              buf[tid].append(v)                        // won the race for v",
        "     barrier",
        "     one thread: offsets = prefix sums of buf sizes; next_size = total   // MERGE",
        "     (implicit barrier)",
        "     every thread: copy buf[tid] into next_frontier[offsets[tid] ...]",
        "     barrier",
        "     one thread: swap(frontier, next_frontier); size = next_size; level++",
        "     (implicit barrier)",
        "  until size == 0",
    ])
    para(doc, "Table 1 lists every shared variable and why it needs no further protection. The key point is that the only data that several threads may write at the same moment is dist[v], and there the atomic compare-and-swap lets exactly one thread succeed.", align="justify")
    table(doc, ["Shared data", "Who writes it", "Why it is safe"], [
        ["dist[v] (std::atomic<int>)", "any thread, during EXPAND", "compare_exchange from −1 to level+1: exactly one thread wins per vertex, so no duplicates and no lost update"],
        ["current frontier array", "nobody during EXPAND", "read-only while the level is expanded"],
        ["next frontier array", "each thread, in its own slice", "slices start at the prefix sums of the buffer sizes, so they never overlap"],
        ["buf[tid] (per-thread buffer)", "its own thread only", "private; aligned to cache lines to avoid false sharing"],
        ["frontier size, level, offsets, next size", "one thread, inside 'single'", "read by all threads only after the implicit barrier of 'single'"],
    ], widths=[1.8, 1.6, 3.1], size=8.5, caption="Table 1. Shared data of the parallel BFS and how each item is protected.")
    para(doc, "**Why relaxed atomics are enough.** The values in dist are used only as visited flags during a level and as the result at the end. The order between different levels is created by the OpenMP barriers, which also make the writes of one level visible to all threads before the next level starts. "
              "A cheap plain read of dist[v] before the compare-and-swap avoids most atomic operations on vertices that are already visited. **Why a barrier per level cannot be avoided:** distances are correct only if level L is complete before level L+1 starts, "
              "so this synchronization is part of the algorithm and not just an implementation detail; it is exactly what makes graphs with many levels expensive (Section 5).", align="justify")

    doc.add_heading("3  Thread configurations and experiment design", level=1)
    table(doc, ["Parameter", "Value"], [
        ["Thread counts", "1, 2 and 4 (the machine has 4 logical CPUs; no over-subscription)"],
        ["Thread binding", f"OMP_PROC_BIND={ENV['OMP_PROC_BIND']}, OMP_PLACES={ENV['OMP_PLACES']}"],
        ["Loop schedule", "dynamic, chunk 64 (set with omp_set_schedule). Static and guided were only tested for correctness, not timed"],
        ["Graphs", "path, 2-D grid, random (average degree 16), R-MAT; 2^14, 2^16, 2^18, 2^20 vertices; density sweep on random graphs, 2^18 vertices, degree 2 to 64"],
        ["Repetitions", "1 warm-up + 7 timed runs per configuration; summary = median (mean, std, CV, min, max also stored)"],
        ["Instrumentation", "one extra run after every timed run measures barrier wait and merge time; these runs never contribute to Execution_Time"],
        ["Experiment type", "strong scaling (fixed problem, more threads). Weak scaling was not measured"],
    ], widths=[1.5, 5.0], size=8.5, caption="Table 2. Experiment parameters.")
    para(doc, "Speedup and efficiency (Student 1’s sequential BFS on the same graph is the baseline), and traversal throughput in TEPS (m_c = undirected edges in the reached component) are", align="justify", after=1)
    equation(doc, EQ_SPEEDUP, eqc)
    equation(doc, EQ_EFF, eqc)
    equation(doc, EQ_TEPS, eqc)
    para(doc, "I report millions of TEPS (MTEPS). To measure the cost of synchronization I recorded, for every level, how long each thread worked in the expand phase; the barrier wait is the slowest thread’s time minus the average thread time, summed over all levels, "
              "and the merge time is the wall time of the merge phase.", align="justify")

    doc.add_heading("4  Results: speedup, efficiency and throughput", level=1)
    rows = []
    for ds, lab in FAM:
        for t in (1, 2, 4):
            rows.append([lab, t] + [f"{sig(sp(ds, n, t))} / {sig(ef(ds, n, t))}" for n in SIZES])
    table(doc, ["Graph", "Threads"] + [f"n = {SZ(n)}" for n in SIZES], rows, widths=[1.4, 0.6, 1.1, 1.1, 1.1, 1.1], size=8,
          caption="Table 3. Speedup / efficiency for all graphs, sizes and thread counts (medians of 7 runs; baseline = sequential BFS on the same graph).")
    figure_row(doc, [("fig2_speedup_vs_threads.png", f"Figure 1. Speedup versus threads for {NMAX:,} vertices."),
                     ("fig5_efficiency_vs_threads.png", f"Figure 2. Efficiency versus threads for {NMAX:,} vertices.")], width=3.1)
    para(doc, f"**Scalability across thread counts.** On the random and R-MAT graphs, speedup grows with the number of threads, to {f2(sp('random_d16', NMAX))} and {f2(sp('rmat_ef8', NMAX))} on four threads at 2^20 vertices, and efficiency stays near {f2(ef('random_d16', NMAX))}–{f2(ef('rmat_ef8', NMAX))}. "
              f"At 2^14 vertices the same graphs reach only {f2(sp('random_d16', NMIN))} and {f2(sp('rmat_ef8', NMIN))} (efficiency {f2(ef('random_d16', NMIN))} and {f2(ef('rmat_ef8', NMIN))}): the whole run lasts less than a millisecond, so thread start-up and barriers take a larger share and the timing noise is high. "
              f"On the grid the speedup stays below 1 ({f2(sp('grid', NMAX, 1))}, {f2(sp('grid', NMAX, 2))} and {f2(sp('grid', NMAX, 4))} for 1, 2 and 4 threads), and on the path it falls with every added thread ({sig(sp('path', NMAX, 1))}, {sig(sp('path', NMAX, 2))} and {sig(sp('path', NMAX, 4))}). "
              f"With one thread the parallel program is already a little faster than the baseline on the random and R-MAT graphs ({f2(sp('random_d16', NMAX, 1))} and {f2(sp('rmat_ef8', NMAX, 1))}); this explains why the two-thread efficiency is slightly above 1 ({f2(ef('random_d16', NMAX, 2))} and {f2(ef('rmat_ef8', NMAX, 2))}), "
              f"because efficiency is measured against the sequential program and not against the one-thread parallel program. The run-to-run CV of these configurations ({min(cv_hi):.0f}–{max(cv_hi):.0f}%) is of similar size, so I read these values as “about 1” and not as true superlinear speedup.", align="justify")
    figure_row(doc, [("fig4_speedup_vs_size.png", f"Figure 3. Speedup with {TMAX} threads versus graph size."),
                     ("fig3_throughput_vs_density.png", f"Figure 4. Throughput versus average degree (random graphs, {NDENS:,} vertices).")], width=3.1)
    rows = []
    for ds, lab in FAM:
        for n in (2 ** 18, NMAX):
            fm = lambda x: f"{x:.2f}" if x < 10 else f"{x:.0f}"
            rows.append([lab, SZ(n), fm(mt(ds, n, 0))] + [fm(mt(ds, n, t)) for t in (1, 2, 4)])
    table(doc, ["Graph", "n", "MTEPS seq.", "MTEPS 1T", "MTEPS 2T", "MTEPS 4T"], rows, widths=[1.5, 0.6, 1.0, 1.0, 1.0, 1.0], size=8.5, caption="Table 4. Traversal throughput in MTEPS (millions of traversed edges per second).")
    para(doc, f"**Throughput.** The best throughput is {mt('rmat_ef8', 2**18, 4):.0f} MTEPS (R-MAT, 2^18 vertices, 4 threads), against {mt('rmat_ef8', 2**18, 0):.0f} MTEPS for the baseline. The path reaches only {mt('path', NMAX, 4):.2f} MTEPS with four threads, against {mt('path', NMAX, 0):.0f} MTEPS sequentially. "
              f"On random graphs with {NDENS:,} vertices the throughput with four threads rises from {f1(dmt(2, TMAX))} MTEPS at average degree 2 to {f1(dmt(64, TMAX))} MTEPS at degree 64 (Figure 4): a denser graph gives each frontier vertex more edges to scan, so there is more work between two barriers. "
              f"The sequential program gains much less ({f1(dmt(2, 0))} to {f1(dmt(64, 0))} MTEPS).", align="justify")

    doc.add_heading("5  Synchronization overhead and performance limits", level=1)
    para(doc, "Every level costs a fixed amount of synchronization (barriers, the prefix-sum step, loop start-up), which does not depend on how much work the level contains. "
              "If a level has little work, this fixed cost dominates. Table 5 gives the cost per level and the measured shares of barrier wait and merge time with four threads; Figure 6 shows the cost per level for all four structures at the largest size.", align="justify")
    rows = []
    for ds, lab in FAM:
        for n in SIZES:
            s = r(ds, n, TMAX)
            rows.append([lab, SZ(n), int(s.Levels), f"{s.Edges_Traversed / s.Levels:,.0f}", f"{100 * s.Sync_Wait_Frac:.1f}", f"{100 * s.Merge_Frac:.1f}", f"{s.Median / s.Levels * 1e6:.2f}"])
    table(doc, ["Graph", "n", "Levels", "Undirected edges per level", "Barrier wait (% of 4T time)", "Merge phase (% of 4T time)", "Time per level, 4T (µs)"], rows, widths=[1.2, 0.5, 0.65, 1.2, 1.05, 1.05, 0.95], size=8,
          caption="Table 5. Work per level and synchronization measurements with 4 threads (shares come from instrumented runs and are approximate).")
    figure_row(doc, [("fig6_overhead_breakdown.png", f"Figure 5. Barrier wait and merge-phase share ({TMAX} threads, {NMAX:,} vertices)."),
                     ("fig10_per_level_cost.png", "Figure 6. Median time per BFS level versus threads (dotted: sequential).")], width=3.1)
    para(doc, f"**The path.** With one vertex per level there is nothing to share, yet each level still executes all steps of the algorithm: the cost per level is {p_lev_us[1]:.2f} µs with one thread, {p_lev_us[2]:.2f} µs with two and {p_lev_us[4]:.2f} µs with four, "
              f"against {seq_level_ns:.1f} ns per vertex sequentially, so the parallel code is {1 / sp('path', NMAX):.0f} times slower with four threads. Even with one thread the cost per level is about {p_lev_us[1] * 1000 / seq_level_ns:.0f} times the sequential cost. "
              "The increase from one to four threads is probably the price of letting more threads meet at every barrier, but I did not test this separately.", align="justify")
    para(doc, f"**The grid.** It has {int(r('grid', NMAX, 0).Levels):,} levels with about {work_per_level['grid']:,.0f} undirected edges each at 2^20 vertices. The sequential program needs {grid_seq_us:.1f} µs per level, the one-thread parallel program {grid_par_us:.1f} µs, "
              f"and with four threads the time per level hardly changes ({r('grid', NMAX, 4).Median / r('grid', NMAX, 4).Levels * 1e6:.1f} µs). The work per level is too small to benefit from more threads, which is why the grid is slower than the baseline at every size. "
              f"**Random and R-MAT graphs.** They have {int(r('random_d16', NMAX, 0).Levels)} and {int(r('rmat_ef8', NMAX, 0).Levels)} levels with more than a million edges each, so the fixed costs are below {100 * max(SYNC4['random_d16'][0] + SYNC4['random_d16'][1], SYNC4['rmat_ef8'][0] + SYNC4['rmat_ef8'][1]):.0f}% of the time and the speedup is close to the thread count.", align="justify")
    para(doc, "**Karp–Flatt analysis.** To express the same idea with one number, I computed the experimentally determined serial fraction of Karp and Flatt from the measured speedup S on p threads:", align="justify", after=1)
    S_ = _sub(_r("S"), _r("p"))
    one_over = lambda x: _frac(_r("1"), x)
    eq_kf = _sub(_r("e"), _r("p")) + _r("=") + _frac(one_over(S_) + _r(MINUS) + one_over(_r("p")), _r("1") + _r(MINUS) + one_over(_r("p")))
    equation(doc, eq_kf, eqc)
    rows = []
    for ds, lab in FAM:
        s2, s4 = exact_speedup(ds, NMAX, 2), exact_speedup(ds, NMAX, 4)
        rows.append([lab, f"{s2:.3f}", f"{karp_flatt(s2, 2):.3f}", f"{s4:.3f}", f"{karp_flatt(s4, 4):.3f}"])
    table(doc, ["Graph (n = 2^20)", "Speedup, 2T", "e (p = 2)", "Speedup, 4T", "e (p = 4)"], rows, widths=[1.7, 1.2, 1.2, 1.2, 1.2], size=8.5,
          caption="Table 6. Karp–Flatt serial fraction e. A value near 0 means almost perfect scaling; a value above 1 means the parallel program is slower than the baseline.")
    para(doc, f"For the random and R-MAT graphs e is small ({karp_flatt(sp('random_d16', NMAX, 4), 4):.3f} and {karp_flatt(sp('rmat_ef8', NMAX, 4), 4):.3f} for four threads), meaning that only a few percent of the work behaves as if it were serial. "
              f"For the grid and the path e is larger than 1, which simply restates that these runs are slower than the baseline. The negative values for two threads occur because the one-thread parallel program is already faster than the sequential baseline, so the speedup on two threads exceeds 2. The Karp–Flatt metric ignores that the baseline and the parallel code are different programs, so I use it only as a compact summary.", align="justify")

    doc.add_heading("6  Limitations and what I would do next", level=1)
    bullets(doc, [
        "Only up to four threads on a shared virtual machine could be tested, so saturation, NUMA effects and many-core behaviour are not visible. The thread count equals the number of logical CPUs, so no over-subscription was studied.",
        f"The CV of short runs is high (up to {cv_max:.0f}%; {n_noisy} of {n_cfg} configurations exceed 15%); speedups of the smallest graphs are therefore uncertain.",
        "The loop schedule was fixed to dynamic,64. Static and guided schedules and other chunk sizes were tested for correctness only; a timing study of the chunk size is open.",
        "Barrier wait and merge time come from separate instrumented runs and are shown as a share of the timed run, so they are approximate; the cost of loop start-up could not be separated.",
        "Only strong scaling was measured, and only synthetic graphs were used. There is no GPU version and no direction-optimizing (bottom-up) variant.",
    ])
    para(doc, "**Next steps (hypotheses, not results).** The measurements suggest running small frontiers sequentially and only switching to parallel expansion when the frontier is large; this should remove most of the penalty on the path and the grid. "
              "For the large frontiers of random and R-MAT graphs, a bottom-up step as in direction-optimizing BFS could reduce the number of edge checks. Both ideas need to be implemented and measured before any claim can be made.", align="justify")
    para(doc, f"**Summary of my contribution.** I implemented and tested the OpenMP BFS, ran {N_PAR_CFG} parallel configurations (21 graphs × 3 thread counts) with 7 timed repetitions each, i.e. {N_PAR_RUNS} timed runs, all verified against the sequential result, "
              "and explained the observed speedups with measured synchronization costs. All numbers in this note come from data/raw/bfs_raw.csv and data/derived/bfs_derived.csv.", align="justify")
    doc.save(OUT_N2 / "Technical_Note_Student2.docx")
