#!/usr/bin/env python3
"""Builds the three supplementary source-code appendices (complete files, syntax-highlighted, with line numbers):

  02_Lucky_Sequential_BFS/Student1_Source_Code_Appendix.docx    bfs_seq.hpp, bfs_seq_main.cpp
  03_Girlfriend_OpenMP_BFS/Student2_Source_Code_Appendix.docx   bfs_par.hpp, bfs_par_main.cpp
  01_Common/Common_Source_Code_Appendix.docx                    graph.hpp, bench_common.hpp, bfs_tests.cpp, Makefile

They are supplementary: the real .cpp/.hpp files are the deliverables and are never replaced by these documents.
Requires: pip install pygments python-docx
"""
from pathlib import Path

from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pygments import lex
from pygments.lexers import CppLexer, MakefileLexer
from pygments.token import Comment, Keyword, Name, Number, String, Token

from build_docs import ROOT, new_doc, para

COMMON, S1, S2 = ROOT, ROOT.parent / "02_Lucky_Sequential_BFS", ROOT.parent / "03_Girlfriend_OpenMP_BFS"

COLORS = [  # (token type, RGB) - first match wins, children of the token type included
    (Comment, (0x00, 0x80, 0x00)),
    (String, (0xA3, 0x15, 0x15)),
    (Comment.Preproc, (0xAF, 0x00, 0xDB)),
    (Keyword.Type, (0x26, 0x7F, 0x99)),
    (Keyword, (0x00, 0x00, 0xFF)),
    (Number, (0x09, 0x86, 0x58)),
    (Name.Function, (0x79, 0x5E, 0x26)),
]


def color_of(tt):
    if tt in Comment.Preproc or tt in Comment.PreprocFile:
        return (0xAF, 0x00, 0xDB)
    for base, rgb in COLORS:
        if tt in base:
            return rgb
    return None


def code_block(doc, path: Path, lexer):
    code = path.read_text(encoding="utf-8")
    lines = [[]]  # list of lines; each line = list of (text, rgb)
    for tt, val in lex(code, lexer):
        rgb = color_of(tt)
        parts = val.split("\n")
        for k, part in enumerate(parts):
            if k > 0:
                lines.append([])
            if part:
                lines[-1].append((part, rgb))
    if lines and not lines[-1]:
        lines.pop()
    width = len(str(len(lines)))
    for no, segs in enumerate(lines, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Pt(0)
        r = p.add_run(f"{no:>{width}}  ")
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        for text, rgb in segs:
            r = p.add_run(text.replace("\t", "    "))
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(7.5)
            if rgb:
                r.font.color.rgb = RGBColor(*rgb)
    return len(lines)


def build(out_path: Path, title: str, who: str, files):
    doc = new_doc(size=10.5)
    para(doc, title, bold=True, align="center", size=15, after=2)
    para(doc, "CS4-S01 Parallel Breadth-First Search — BCSE412L Parallel Computing, VIT Vellore", align="center", size=10, after=2)
    para(doc, who, align="center", size=10, after=8)
    para(doc, "Supplementary appendix. It reproduces the complete source files listed below, unchanged and with line numbers, for reading and printing. "
              "The actual compilable files are delivered separately in the project folders and are the authoritative versions. "
              "Build and run instructions are in 01_Common/README.md.", align="justify", size=10)
    toc = ", ".join(f.name for f, _ in files)
    para(doc, f"**Files in this appendix:** {toc}", size=10)
    for i, (f, lexer) in enumerate(files):
        h = doc.add_heading(f"{f.name}", level=1)
        rel = f.relative_to(ROOT.parent)
        para(doc, f"Location: {rel.as_posix()}", italic=True, size=9, after=3)
        n = code_block(doc, f, lexer)
        print(f"  {f.name}: {n} lines")
        if i < len(files) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    cpp, mk = CppLexer(), MakefileLexer()
    build(S1 / "Student1_Source_Code_Appendix.docx", "Student 1 — Source Code Appendix", "Sequential BFS baseline: [Student 1 name, Reg. No.]",
          [(S1 / "bfs_seq.hpp", cpp), (S1 / "bfs_seq_main.cpp", cpp)])
    build(S2 / "Student2_Source_Code_Appendix.docx", "Student 2 — Source Code Appendix", "OpenMP parallel BFS: [Student 2 name, Reg. No.]",
          [(S2 / "bfs_par.hpp", cpp), (S2 / "bfs_par_main.cpp", cpp)])
    build(COMMON / "Common_Source_Code_Appendix.docx", "Common Source Code Appendix", "Shared by both students",
          [(COMMON / "graph.hpp", cpp), (COMMON / "bench_common.hpp", cpp), (COMMON / "bfs_tests.cpp", cpp), (COMMON / "Makefile", mk)])
